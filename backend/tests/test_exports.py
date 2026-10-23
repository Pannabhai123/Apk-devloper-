from io import BytesIO

import fitz
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.schemas import GenerateRequest
from app.services.pdf_export import generate_pdf
from app.services.pptx_service import generate_pptx


def sample_questions():
    return [
        {
            "id": "q1", "number": "1", "text": "भारत की राजधानी क्या है?",
            "options": [{"label": "A", "text": "मुंबई"}, {"label": "B", "text": "नई दिल्ली"}, {"label": "C", "text": "कोलकाता"}, {"label": "D", "text": "चेन्नई"}],
            "correct_answer": "B", "answer_text": "नई दिल्ली", "explanation": "नई दिल्ली भारत की राजधानी है।",
            "answer_source": "pdf", "answer_confidence": "certain", "page_number": 1,
        },
        {
            "id": "q2", "number": "2", "text": "Which planet is known as the Red Planet?",
            "options": [{"label": "A", "text": "Earth"}, {"label": "B", "text": "Venus"}, {"label": "C", "text": "Mars"}, {"label": "D", "text": "Jupiter"}],
            "correct_answer": "C", "explanation": "Iron oxide gives Mars its colour.",
            "answer_source": "ai", "answer_confidence": "uncertain", "page_number": 2,
        },
    ]


@pytest.mark.parametrize("mode,expected", [("A", 2), ("B", 4), ("C", 4), ("D", 4)])
def test_pptx_is_editable_widescreen_and_reopens(tmp_path, mode, expected):
    output = tmp_path / f"mode_{mode}.pptx"
    settings = GenerateRequest(mode=mode).model_dump()
    count, _ = generate_pptx(sample_questions(), settings, output)
    assert count == expected
    prs = Presentation(str(output))
    assert len(prs.slides) == expected
    assert round(prs.slide_width / prs.slide_height, 2) == 1.78
    all_text = "\n".join(
        shape.text for slide in prs.slides for shape in slide.shapes
        if hasattr(shape, "text_frame") and shape.has_text_frame
    )
    assert "भारत की राजधानी क्या है?" in all_text
    assert "नई दिल्ली" in all_text
    if mode != "A":
        assert "AI-GENERATED" in all_text
        assert "UNCERTAIN" in all_text
    assert not any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for slide in prs.slides for shape in slide.shapes
    )


def test_long_question_uses_continuation_instead_of_clipping(tmp_path):
    questions = [dict(sample_questions()[0], text="Long Hindi प्रश्न " * 100)]
    count, _ = generate_pptx(questions, GenerateRequest(mode="A").model_dump(), tmp_path / "long.pptx")
    assert count > 1


def test_long_options_are_paginated_in_pptx_and_pdf(tmp_path):
    question = dict(sample_questions()[0])
    question["options"] = [
        {"label": "A", "text": "Long option विकल्प " * 180},
        {"label": "B", "text": "नई दिल्ली"},
        {"label": "C", "text": "कोलकाता"},
        {"label": "D", "text": "चेन्नई"},
    ]
    settings = GenerateRequest(mode="D", custom_answer_slide=True, custom_show_options_on_answer=True).model_dump()
    slide_count, plans = generate_pptx([question], settings, tmp_path / "long_options.pptx")
    assert slide_count > 2
    assert any("cont." in option["label"] for plan in plans for option in plan["options"])
    assert all(len(option["text"]) <= 450 for plan in plans for option in plan["options"])
    pdf_count = generate_pdf([question], settings, tmp_path / "long_options.pdf")
    assert pdf_count == slide_count
    with fitz.open(tmp_path / "long_options.pdf") as document:
        assert len(document) == slide_count


def test_pdf_export_opens_with_every_slide(tmp_path):
    output = tmp_path / "slides.pdf"
    count = generate_pdf(sample_questions(), GenerateRequest(mode="C").model_dump(), output)
    assert count == 4
    assert output.read_bytes().startswith(b"%PDF")
    with fitz.open(output) as document:
        assert len(document) == 4
        assert document[0].rect.width > document[0].rect.height
