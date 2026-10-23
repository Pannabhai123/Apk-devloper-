from io import BytesIO
from pathlib import Path

import fitz
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app

SAMPLE = Path(__file__).resolve().parents[2] / "samples" / "Dream_Classes_Bilingual_Sample.pdf"


def test_complete_pdf_to_ppt_workflow():
    with TestClient(app) as client:
        health = client.get("/api/health")
        assert health.status_code == 200
        assert health.json()["capabilities"]["pptx"] is True
        assert "api_key" not in health.text.lower()

        with SAMPLE.open("rb") as source:
            uploaded = client.post("/api/projects", files={"file": (SAMPLE.name, source, "application/pdf")})
        assert uploaded.status_code == 202, uploaded.text
        project_id = uploaded.json()["id"]

        # FastAPI executes BackgroundTasks before TestClient finishes the request.
        project = client.get(f"/api/projects/{project_id}").json()
        assert project["status"] == "ready", project
        assert project["total_pages"] == 2
        assert project["processed_pages"] == 2
        assert project["readable_pages"] == 2
        assert project["question_count"] == 6

        response = client.get(f"/api/projects/{project_id}/questions")
        questions = response.json()
        assert len(questions) == 6
        assert questions[3]["text"] == "भारत की राजधानी क्या है?"
        assert questions[3]["correct_answer"] == "B"
        assert questions[3]["explanation"] == "नई दिल्ली भारत की राजधानी है।"
        missing = next(q for q in questions if not q["correct_answer"])

        # No key is configured: the app must fail honestly without changing the question.
        ai_response = client.post(f"/api/projects/{project_id}/generate-answers")
        assert ai_response.status_code == 503
        assert "AI_API_KEY" in ai_response.json()["detail"]

        edited = client.patch(f"/api/questions/{missing['id']}", json={
            "correct_answer": "A", "explanation": "Oxygen is used in respiration.",
        })
        assert edited.status_code == 200
        assert edited.json()["answer_source"] == "manual"
        assert "missing_answer" not in edited.json()["flags"]

        generated = client.post(f"/api/projects/{project_id}/generate", json={"mode": "C"})
        assert generated.status_code == 200, generated.text
        assert generated.json()["slide_count"] == 12

        pptx_response = client.get(f"/api/projects/{project_id}/download/pptx")
        assert pptx_response.status_code == 200
        presentation = Presentation(BytesIO(pptx_response.content))
        assert len(presentation.slides) == 12
        text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
        assert "भारत की राजधानी क्या है?" in text
        assert "Oxygen is used in respiration." in text

        pdf_response = client.get(f"/api/projects/{project_id}/download/pdf")
        assert pdf_response.status_code == 200, pdf_response.text
        with fitz.open(stream=pdf_response.content, filetype="pdf") as document:
            assert len(document) == 12

        questions_download = client.get(f"/api/projects/{project_id}/download/questions")
        assert questions_download.status_code == 200
        assert len(questions_download.json()["questions"]) == 6
        answer_key = client.get(f"/api/projects/{project_id}/download/answer-key")
        assert answer_key.status_code == 200
        assert "Oxygen is used in respiration." in answer_key.content.decode("utf-8")
        assert len(client.get("/api/exports").json()) == 2

        assert client.delete(f"/api/projects/{project_id}").status_code == 204
        assert client.get(f"/api/projects/{project_id}").status_code == 404


def test_unreadable_pdf_and_invalid_upload_report_errors():
    with TestClient(app) as client:
        invalid = client.post("/api/projects", files={"file": ("fake.pdf", b"not a pdf", "application/pdf")})
        assert invalid.status_code == 400

        blank_path = Path("/tmp/dream_classes_blank.pdf")
        document = fitz.open()
        document.new_page()
        document.save(blank_path)
        document.close()
        with blank_path.open("rb") as source:
            response = client.post("/api/projects", files={"file": ("blank.pdf", source, "application/pdf")})
        assert response.status_code == 202
        project_id = response.json()["id"]
        project = client.get(f"/api/projects/{project_id}").json()
        assert project["status"] == "failed"
        assert project["total_pages"] == 1
        assert project["processed_pages"] == 1
        assert project["unreadable_pages"][0]["page"] == 1
        assert client.delete(f"/api/projects/{project_id}").status_code == 204
