"""Genuine editable 16:9 PowerPoint renderer using python-pptx."""
from math import ceil
from pathlib import Path
import unicodedata
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from .slide_service import build_slide_plan, enrich_answer_options

SLIDE_W = 13.333333
SLIDE_H = 7.5


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


def _visual_length(text: str) -> int:
    return sum(
        0 if unicodedata.combining(char) else (2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1)
        for char in text
    )


def _estimated_lines(text: str, width: float, font_size: float) -> int:
    capacity = max(8, int(width * 72 / (font_size * 0.53)))
    return max(1, sum(max(1, ceil(_visual_length(line) / capacity)) for line in text.splitlines() or [""]))


def _fit_size(items: list[tuple[str, float, float]], preferred: int, minimum: int = 13) -> int:
    """items are (text, box width inches, box height inches)."""
    size = preferred
    while size > minimum:
        overflow = False
        for text, width, height in items:
            lines = _estimated_lines(text, width, size)
            if lines * size * 1.23 / 72 > height:
                overflow = True
                break
        if not overflow:
            break
        size -= 1
    return size


def _textbox(slide, text: str, x: float, y: float, w: float, h: float, *,
             size: int, color: str, font: str, bold: bool = False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin: float = 0.03, fit: bool = True):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    if fit:
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = 1.08
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return shape


def _rect(slide, x: float, y: float, w: float, h: float, fill: str,
          line: str | None = None, radius: bool = False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def _base_slide(prs: Presentation, settings: dict[str, Any], plan: dict[str, Any], index: int, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = _rgb(settings["background_color"])
    _textbox(slide, "DREAM CLASSES", 0.65, 0.28, 3.4, 0.34,
             size=12, color=settings["highlight_color"], font=settings["font_family"], bold=True)
    _textbox(slide, "PPT MAKER AI", 10.25, 0.28, 2.4, 0.34,
             size=9, color="92A59A", font=settings["font_family"], bold=True, align=PP_ALIGN.RIGHT)
    _rect(slide, 0.65, 0.73, 12.03, 0.012, "2A3A31")
    _textbox(slide, "DREAM SIR  •  KOTHWARA, DOBHI, GAYA", 0.65, 7.12, 7.0, 0.2,
             size=7, color="71877A", font=settings["font_family"], fit=False)
    _textbox(slide, f"{index:02d} / {total:02d}", 11.25, 7.09, 1.4, 0.23,
             size=8, color="71877A", font=settings["font_family"], align=PP_ALIGN.RIGHT, fit=False)
    return slide


def _question_slide(prs, settings, plan, index, total):
    slide = _base_slide(prs, settings, plan, index, total)
    font = settings["font_family"]
    highlight = settings["highlight_color"]
    label = f"Q{plan['number']}" if settings.get("show_question_numbers", True) else "QUESTION"
    if plan.get("continuation"):
        label += " · CONTINUED"
    _textbox(slide, label.upper(), 0.67, 0.94, 4.2, 0.4,
             size=13, color=highlight, font=font, bold=True)
    _textbox(slide, f"SOURCE PAGE {plan['page_number']}", 9.8, 0.96, 2.85, 0.34,
             size=9, color="92A59A", font=font, align=PP_ALIGN.RIGHT)

    has_options = bool(plan["options"])
    include_answer = plan.get("include_answer")
    question_height = 1.72 if has_options else (3.6 if include_answer else 5.25)
    qsize = _fit_size(
        [(plan["text"], 11.95, question_height)], settings["question_font_size"], 14
    )
    _textbox(slide, plan["text"], 0.67, 1.39, 11.95, question_height,
             size=qsize, color=settings["question_color"], font=font, bold=True,
             valign=MSO_ANCHOR.MIDDLE)

    if has_options:
        options = plan["options"]
        compact = len(options) > 2 and max(_visual_length(o["text"]) for o in options) < 95 and len(plan["text"]) < 520
        bottom = 5.92 if include_answer else 6.82
        start_y = 3.34
        available = bottom - start_y
        if compact:
            rows = ceil(len(options) / 2)
            box_h = min(1.18, (available - (rows - 1) * 0.16) / rows)
            for option_index, option in enumerate(options):
                row, col = divmod(option_index, 2)
                x = 0.67 + col * 6.08
                y = start_y + row * (box_h + 0.16)
                w = 5.87
                _rect(slide, x, y, w, box_h, "101D16", "263C30", radius=True)
                text = f"{option['label']}.  {option['text']}"
                osize = _fit_size([(text, w - 0.3, box_h - 0.1)], settings["option_font_size"], 13)
                _textbox(slide, text, x + 0.15, y + 0.05, w - 0.3, box_h - 0.1,
                         size=osize, color=settings["option_color"], font=font,
                         valign=MSO_ANCHOR.MIDDLE)
        else:
            count = max(1, len(options))
            gap = 0.12
            box_h = min(0.82, (available - (count - 1) * gap) / count)
            for option_index, option in enumerate(options):
                y = start_y + option_index * (box_h + gap)
                _rect(slide, 0.67, y, 11.95, box_h, "101D16", "263C30", radius=True)
                text = f"{option['label']}.  {option['text']}"
                osize = _fit_size([(text, 11.62, box_h - 0.08)], settings["option_font_size"], 12)
                _textbox(slide, text, 0.83, y + 0.04, 11.62, box_h - 0.08,
                         size=osize, color=settings["option_color"], font=font,
                         valign=MSO_ANCHOR.MIDDLE)
    if include_answer:
        answer = plan.get("correct_answer") or "Answer not provided"
        answer_text = plan.get("answer_text") or ""
        content = f"ANSWER: {answer}" + (f" — {answer_text}" if answer_text and answer_text != answer else "")
        if plan.get("explanation"):
            content += f"\nExplanation: {plan['explanation']}"
        _textbox(slide, content, 0.75, 5.98, 11.75, 0.82,
                 size=_fit_size([(content, 11.75, 0.82)], 17, 11),
                 color=highlight, font=font, bold=True, valign=MSO_ANCHOR.MIDDLE)
    if plan.get("continues"):
        _textbox(slide, "CONTINUED ON NEXT SLIDE →", 8.7, 6.8, 3.9, 0.2,
                 size=8, color=highlight, font=font, align=PP_ALIGN.RIGHT)
    return slide


def _answer_slide(prs, settings, plan, index, total):
    slide = _base_slide(prs, settings, plan, index, total)
    font = settings["font_family"]
    highlight = settings["highlight_color"]
    answer = plan.get("correct_answer")
    _textbox(slide, f"Q{plan['number']}  /  ANSWER REVEAL", 0.67, 0.97, 5.5, 0.4,
             size=13, color=highlight, font=font, bold=True)
    if plan["answer_source"] == "ai":
        badge = "AI-GENERATED  •  VERIFY BEFORE TEACHING"
        if plan["answer_confidence"] == "uncertain":
            badge = "AI-GENERATED  •  UNCERTAIN  •  VERIFY"
        _rect(slide, 8.52, 0.92, 4.1, 0.43, "302015", "694323", radius=True)
        _textbox(slide, badge, 8.64, 1.01, 3.86, 0.21,
                 size=7, color=highlight, font=font, bold=True, align=PP_ALIGN.CENTER, fit=False)

    if not answer:
        _textbox(slide, "Answer not provided", 0.75, 2.15, 11.75, 0.9,
                 size=34, color=highlight, font=font, bold=True, align=PP_ALIGN.CENTER)
        _textbox(slide, "Review this question or use configured AI answer generation.", 1.2, 3.25, 10.9, 0.55,
                 size=17, color="B5C2BA", font=font, align=PP_ALIGN.CENTER)
    else:
        _textbox(slide, "CORRECT OPTION", 0.75, 1.72, 3.0, 0.35,
                 size=11, color="9BACA1", font=font, bold=True)
        _textbox(slide, str(answer), 0.72, 2.04, 2.1, 1.4,
                 size=66, color=highlight, font=font, bold=True, valign=MSO_ANCHOR.MIDDLE)
        answer_text = plan.get("answer_text", "")
        if answer_text:
            _textbox(slide, answer_text, 2.75, 2.05, 9.55, 1.35,
                     size=_fit_size([(answer_text, 9.55, 1.35)], 29, 15),
                     color=settings["question_color"], font=font, bold=True, valign=MSO_ANCHOR.MIDDLE)

    explanation = plan.get("explanation")
    has_options = bool(plan.get("options"))
    if explanation:
        box_height = 1.72 if has_options else 2.65
        text_height = 0.82 if has_options else 1.65
        _rect(slide, 0.7, 3.85, 11.92, box_height, "101D16", "2A3D31", radius=True)
        _textbox(slide, "EXPLANATION", 0.98, 4.08, 3.2, 0.35,
                 size=11, color=highlight, font=font, bold=True)
        _textbox(slide, explanation, 0.98, 4.53, 11.35, text_height,
                 size=_fit_size([(explanation, 11.35, text_height)], 20, 12),
                 color="F2F5F2", font=font)
    elif answer:
        text_height = 1.3 if has_options else 1.75
        _textbox(slide, plan["text"], 1.0, 4.15, 11.3, text_height,
                 size=_fit_size([(plan["text"], 11.3, text_height)], 17, 12),
                 color="A9B8AE", font=font, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if has_options:
        # Mode D can request options on answer slides; PowerPoint autofit remains enabled as a safeguard.
        option_text = "    ".join(f"{o['label']}. {o['text']}" for o in plan["options"])
        _textbox(slide, option_text, 0.9, 5.82, 11.5, 0.68,
                 size=_fit_size([(option_text, 11.5, 0.68)], 14, 10),
                 color=settings["option_color"], font=font, align=PP_ALIGN.CENTER)
    return slide


def generate_pptx(questions: list[dict[str, Any]], settings: dict[str, Any], output_path: str | Path,
                  title: str = "Dream Classes Questions") -> tuple[int, list[dict[str, Any]]]:
    if not questions:
        raise ValueError("Add at least one extracted question before generating a presentation.")
    plans = build_slide_plan(questions, settings)
    enrich_answer_options(plans, questions)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = title
    prs.core_properties.author = "Dream Sir · Dream Classes"
    prs.core_properties.subject = "Editable PDF-to-PowerPoint question presentation"
    prs.core_properties.comments = "Generated by Dream Classes PPT Maker AI"
    for index, plan in enumerate(plans, start=1):
        if plan["kind"] == "question":
            _question_slide(prs, settings, plan, index, len(plans))
        else:
            _answer_slide(prs, settings, plan, index, len(plans))
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    # Re-open the actual file as an integrity check before making it downloadable.
    verified = Presentation(str(output_path))
    if len(verified.slides) != len(plans):
        raise RuntimeError("PPTX validation failed: slide count changed after saving.")
    return len(plans), plans
